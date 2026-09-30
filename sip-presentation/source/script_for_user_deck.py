"""Align the speaker notes of the presenter's own 35-slide deck with its slide order, then export
them as JSON for the script document (script_docx.js).

usage: python3 script_for_user_deck.py <in.pptx> <out.pptx> <notes.json>
"""
import sys, json, re, copy
from pptx import Presentation

HEADS = ["WHAT I SAY", "WHY IT MATTERS", "WHAT THE PANEL SHOULD NOTICE", "LIKELY PANEL QUESTIONS & MY ANSWERS", "TRANSITION TO NEXT SLIDE", "IMAGE SOURCES"]


def parse(t):
    parts, cur = {}, None
    for line in t.split("\n"):
        if line.strip() in HEADS:
            cur = line.strip(); parts[cur] = []
        elif cur:
            parts[cur].append(line)
    return {k: "\n".join(v).strip() for k, v in parts.items()}


def render(p):
    return "\n\n".join(f"{h}\n{p[h]}" for h in HEADS if p.get(h))


LANDING = {
    "WHAT I SAY": "Good morning. I'm Shreejita Srivastava, and I'm presenting my Summer Internship Project, carried out in the Operations Department at JHS Svendgaard Laboratories, Kala-Amb, from the first of June to the seventeenth of July 2026. My internal guide was Dr. Vineet Tiwari and my external guide was Mr. Paramveer Singh. The project was the development of an Excel-based Manufacturing Execution System. I'll take you through it the way it actually happened: the plant first, then the problem it showed me, then what I built.",
    "WHY IT MATTERS": "It sets the context in one breath: where, when, with whom, and what. It also tells the panel the presentation is a story from the shop floor, not a software demo.",
    "WHAT THE PANEL SHOULD NOTICE": "Both institutions on one page, and both guides named.",
    "LIKELY PANEL QUESTIONS & MY ANSWERS": "Q1. Why an Excel-based MES?\nA: Because it's what the plant already uses: no new licences, no IT approval, no training burden. It let me build and test a working prototype within seven weeks. The data model carries over to a database later.",
    "TRANSITION TO NEXT SLIDE": "Let me start with where I was.",
}

TITLES = ["Landing page", "Where was I?", "What does it look like on the ground?", "What did the plant teach me about the system?",
          "So I asked myself: what would the system need to answer?", "Imagine a customer sends us this", "The complaint",
          "How do we find out what happened?", "What does the pack tell us?", "Why isn't the Batch Number enough?",
          "What identifies one manufacturing occurrence?", "How does material move through the plant?", "The plant had the data",
          "Where was the information, and why was it slow to use?", "What did twelve hours on one line show me?",
          "So where did 27,500 units go?", "How did the project evolve?", "How did I structure it?", "What connects the records?",
          "Before I show you the system... (video)", "Can management see the overall picture?", "Can we drill from a KPI to the records behind it?",
          "Which manufacturing occurrence are we investigating?", "One record. One key.", "Once we find the right record, what can we see?",
          "For loose bristles, where would QA look first?", "Can we investigate the complaint end to end?",
          "If a material lot is implicated, where else did it go?", "But what tells us how efficiently it was produced?",
          "How efficiently did the lines perform?", "What changed?", "What did the project teach me?",
          "The hardest part wasn't building the workbook", "It was understanding the operation...", "Thank you"]

# slide number (in this deck) -> {section: (old, new) or new text}
EDITS = {
    2: {"TRANSITION TO NEXT SLIDE": "Here's what that looks like on the ground."},
    3: {"TRANSITION TO NEXT SLIDE": "During the induction I spent time in the paste, moulding and tufting plants, and each one taught me something the system would later have to do."},
    4: {"TRANSITION TO NEXT SLIDE": "Those lessons turned into a set of questions."},
    5: {"WHAT I SAY": ("And can I calculate OEE to explain the gap I saw on the line?", "And can I calculate OEE to explain where production time is lost?"),
        "WHAT THE PANEL SHOULD NOTICE": "Questions 1 to 4 are about traceability; I'll show you why with a complaint scenario next. Questions 5 and 6 are about performance; they come from the line study I'll show later.",
        "TRANSITION TO NEXT SLIDE": "Let me show you why the first question matters so much."},
    6: {"WHAT I SAY": ("Good morning. Before I tell you about my internship, I want to start with a situation. Imagine a customer sends us a complaint about one of our products. That's where this whole project begins.",
                       "Let me show you why those questions matter, with a situation. Imagine a customer sends us a complaint about one of our products. That's where the traceability side of this project begins."),
        "WHY IT MATTERS": "It turns the requirements into a concrete, human situation that the rest of the deck answers."},
    11: {"TRANSITION TO NEXT SLIDE": "To solve this, I had to go back to where the information is created: how material moves through the plant."},
    12: {"TRANSITION TO NEXT SLIDE": "And that's what led me to the real finding."},
    16: {"TRANSITION TO NEXT SLIDE": "At this point I had two problems, a traceability problem and a visibility problem, which is exactly what those six questions were about. Here's how the project evolved from there."},
    32: {"TRANSITION TO NEXT SLIDE": "If I had to sum up the whole internship in one line..."},
    35: {"LIKELY PANEL QUESTIONS & MY ANSWERS": "Q1. What are the limitations?\nA: It's a working prototype, not an enterprise MES. Excel will slow down as registers grow; it isn't built for many users editing at once; there are no user roles or audit trail; it depends on disciplined manual entry; and there's no machine data capture, so OEE is calculated after the event. It also runs on a demonstration dataset, not the company's records.\n\nQ2. What is the future scope?\nA: Barcode scanning to capture batch and lot data at source, RFID for movement tracking, SQL migration for scale, concurrency, roles and audit, IoT for run time and counts from equipment, ERP integration, Power BI for wider reporting, and predictive analytics once a longer verified history exists. The sensible order is data discipline, then barcode capture, then database migration.\n\nQ3. What was your personal contribution?\nA: The process study and line observation, the requirement definition, the data model including the Production Run ID and four-field identification, building all 27 sheets, and the testing: 106 calculation checks and 80 filter-context checks."},
}


def main(src, dst, out_json):
    prs = Presentation(src)
    export = []
    for i, s in enumerate(prs.slides, 1):
        tf = s.notes_slide.notes_text_frame
        if tf is None:  # notes page without a body placeholder (the landing page): borrow one
            donor = next(x.notes_slide for x in prs.slides if x.has_notes_slide and x.notes_slide.notes_placeholder is not None)
            body = copy.deepcopy(donor.notes_placeholder._element)
            s.notes_slide.shapes._spTree.append(body)
            tf = s.notes_slide.notes_text_frame
            tf.text = ""
        p = parse(tf.text) if tf.text.strip() else dict(LANDING)
        for sec, v in EDITS.get(i, {}).items():
            if isinstance(v, tuple):
                assert v[0] in p[sec], (i, sec, v[0][:50]); p[sec] = p[sec].replace(v[0], v[1])
            else:
                p[sec] = v
        tf.text = render(p)
        export.append({"n": i, "title": TITLES[i - 1], **p})
    prs.save(dst)
    json.dump(export, open(out_json, "w"), indent=1, ensure_ascii=False)
    print("slides:", len(export))


if __name__ == "__main__":
    main(*sys.argv[1:4])
